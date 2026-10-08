import json
import secrets
from io import BytesIO
from datetime import datetime
from functools import wraps

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from django.conf import settings
from django.contrib.auth.hashers import make_password
from django.core.files.storage import default_storage
from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views import View
from django.shortcuts import redirect, render
from django.utils.decorators import method_decorator
from django.utils.text import get_valid_filename
from mongoengine.errors import NotUniqueError, ValidationError

from .models import Usuario, TokenAcesso, Alerta, Camera, Relatorio, PerfilCrianca


PERFIL_CRIANCA_CAMPOS = {
    'nome_completo': 'nome_completo',
    'nome_social': 'nome_social',
    'data_nascimento': 'data_nascimento',
    'sexo': 'sexo',
    'foto_perfil': 'foto_perfil_url',
    'diagnostico_principal': 'diagnostico_principal',
    'nivel_suporte': 'nivel_suporte',
    'data_diagnostico': 'data_diagnostico',
    'laudos_relatorios': 'laudo_url',
    'condicoes_saude': 'condicoes',
    'alergias': 'alergias',
    'medicamentos': 'medicacoes',
    'frequencia_crises': 'frequencia_crises',
    'atividades_favoritas': 'atividades_favoritas',
    'principais_gatilhos': 'gatilhos',
    'estrategias_acalmar': 'o_que_ajuda_acalmar',
    'forma_comunicacao': 'comunicacao',
    'horario_acordar': 'horario_acordar',
    'horario_dormir': 'horario_dormir',
    'alimentacao': 'alimentacao',
    'atividades_terapeuticas_escolares': 'atividades',
    'observacoes': 'observacoes',
    'contato_emergencia': 'contato_emergencia',
    'relacao_com_crianca': 'relacao_com_crianca',
    'foto_frente': 'foto_frente_url',
    'foto_lado_direito': 'foto_direita_url',
    'foto_lado_esquerdo': 'foto_esquerda_url',
    'foto_cima': 'foto_cima_url',
    'foto_baixo': 'foto_baixo_url',
    'foto_sorrindo': 'foto_sorrindo_url',
}


def lista_do_formulario(valor):
    return [item.strip() for item in valor.split(',') if item.strip()]


def data_do_formulario(valor):
    if not valor:
        return None


def salvar_foto(arquivo):
    if arquivo is None:
        return None
    nome = get_valid_filename(arquivo.name)
    return default_storage.save(f'perfis_criancas/{secrets.token_hex(8)}_{nome}', arquivo)
    try:
        return datetime.strptime(valor, '%Y-%m-%d')
    except ValueError:
        return None


def perfil_crianca_do_formulario(dados, arquivos=None, perfil_atual=None):
    nome = dados.get('crianca_nome_completo', '').strip()
    fotos_enviadas = arquivos and any(arquivos.get(f'crianca_{campo}') for campo in {
        'foto_perfil', 'foto_frente', 'foto_lado_direito', 'foto_lado_esquerdo',
        'foto_cima', 'foto_baixo', 'foto_sorrindo',
    })
    preenchido = nome or fotos_enviadas or any(dados.get(f'crianca_{campo}', '').strip() for campo in PERFIL_CRIANCA_CAMPOS if campo != 'nome_completo')
    if not preenchido:
        return None
    if not nome:
        raise ValueError('O nome completo da criança é obrigatório quando houver dados de perfil.')

    perfil = {'nome_completo': nome}
    for campo, campo_modelo in PERFIL_CRIANCA_CAMPOS.items():
        if campo == 'nome_completo':
            continue
        valor = dados.get(f'crianca_{campo}', '').strip()
        arquivo = arquivos.get(f'crianca_{campo}') if arquivos else None
        if campo in {'foto_perfil', 'foto_frente', 'foto_lado_direito', 'foto_lado_esquerdo', 'foto_cima', 'foto_baixo', 'foto_sorrindo'}:
            perfil[campo_modelo] = salvar_foto(arquivo) if arquivo else getattr(perfil_atual, campo_modelo, None)
        elif campo in {'data_nascimento', 'data_diagnostico'}:
            perfil[campo_modelo] = data_do_formulario(valor)
        elif campo in {'condicoes_saude', 'alergias', 'atividades_favoritas', 'principais_gatilhos', 'estrategias_acalmar', 'alimentacao', 'atividades_terapeuticas_escolares'}:
            perfil[campo_modelo] = lista_do_formulario(valor)
        elif campo == 'medicamentos':
            if not valor:
                perfil[campo_modelo] = []
            else:
                try:
                    medicamentos = json.loads(valor)
                except json.JSONDecodeError:
                    medicamentos = lista_do_formulario(valor)
                perfil[campo_modelo] = medicamentos if isinstance(medicamentos, list) else []
        else:
            perfil[campo_modelo] = valor or None

    perfil['usa_medicamentos'] = dados.get('crianca_usa_medicamentos') in {'on', 'true', '1'}
    perfil['possui_crises_epilepticas'] = dados.get('crianca_possui_crises_epilepticas') in {'on', 'true', '1'}
    perfil['possui_sonecas'] = dados.get('crianca_faz_soneca') in {'on', 'true', '1'}
    perfil['seletividade_alimentar'] = dados.get('crianca_possui_seletividade_alimentar') in {'on', 'true', '1'}
    return PerfilCrianca(**perfil)


def serializar_usuario(usuario):
    perfil = None
    if usuario.perfil_crianca is not None:
        perfil = json.loads(usuario.perfil_crianca.to_json())
    return {
        'id': str(usuario.id), 'nome_completo': usuario.nome_completo,
        'email': usuario.email, 'tipo_usuario': usuario.tipo_usuario,
        'ip_camera': usuario.ip_camera, 'perfil_crianca': perfil,
    }


def serializar_perfil_crianca(perfil):
    if perfil is None:
        return None
    dados = json.loads(perfil.to_json())
    return {
        'informacoes_basicas': {
            'nome_completo': dados.get('nome_completo'),
            'nome_social': dados.get('nome_social'),
            'data_nascimento': dados.get('data_nascimento'),
            'sexo': dados.get('sexo'),
            'foto_perfil': dados.get('foto_perfil_url'),
        },
        'diagnostico': {
            'diagnostico_principal': dados.get('diagnostico_principal'),
            'nivel_suporte': dados.get('nivel_suporte'),
            'data_diagnostico': dados.get('data_diagnostico'),
            'laudos_relatorios': dados.get('laudo_url'),
        },
        'saude': {
            'condicoes_saude': dados.get('condicoes'),
            'alergias': dados.get('alergias'),
            'usa_medicamentos': dados.get('usa_medicamentos', bool(dados.get('medicacoes'))),
            'medicamentos': dados.get('medicacoes'),
            'possui_crises_epilepticas': dados.get('possui_crises_epilepticas'),
            'frequencia_crises': dados.get('frequencia_crises'),
        },
        'preferencias_comportamentos': {
            'atividades_favoritas': dados.get('atividades_favoritas'),
            'principais_gatilhos': dados.get('gatilhos'),
            'estrategias_acalmar': dados.get('o_que_ajuda_acalmar'),
            'forma_comunicacao': dados.get('comunicacao'),
        },
        'rotina_diaria': {
            'horario_acordar': dados.get('horario_acordar'),
            'horario_dormir': dados.get('horario_dormir'),
            'faz_soneca': dados.get('possui_sonecas'),
            'alimentacao': dados.get('alimentacao'),
            'possui_seletividade_alimentar': dados.get('seletividade_alimentar'),
            'atividades_terapeuticas_escolares': dados.get('atividades'),
        },
        'informacoes_adicionais': {
            'observacoes': dados.get('observacoes'),
            'contato_emergencia': dados.get('contato_emergencia'),
            'relacao_com_crianca': dados.get('relacao_com_crianca'),
        },
        'reconhecimento_facial': {
            'foto_frente': dados.get('foto_frente_url'),
            'foto_lado_direito': dados.get('foto_direita_url'),
            'foto_lado_esquerdo': dados.get('foto_esquerda_url'),
            'foto_cima': dados.get('foto_cima_url'),
            'foto_baixo': dados.get('foto_baixo_url'),
            'foto_sorrindo': dados.get('foto_sorrindo_url'),
        },
    }


def serializar_rastreamento(usuario):
    camera = Camera.objects(responsavel=usuario).first()
    if camera is None:
        return {'camera': None, 'perfil_carregado': usuario.perfil_crianca is not None}
    return {
        'camera': {
            'id': str(camera.id),
            'endereco_ip': camera.endereco_ip,
            'nome_camara': camera.nome_camara,
        },
        'perfil_carregado': usuario.perfil_crianca is not None,
        'websocket': f'/ws/camera/{camera.id}/',
    }


def contexto_admin(admin_usuario=None):
    usuarios = list(Usuario.objects.order_by('-criado_em'))
    cameras = list(Camera.objects.order_by('-id'))
    alertas = list(Alerta.objects.order_by('-timestamp'))
    tokens = list(TokenAcesso.objects)
    relatorios = list(Relatorio.objects.order_by('-data_geracao'))
    perfis_completos = sum(1 for usuario in usuarios if usuario.perfil_crianca is not None)
    receita = sum((usuario.valor or 0) for usuario in usuarios)
    mes_atual = datetime.utcnow().month
    ano_atual = datetime.utcnow().year
    novos_usuarios = sum(
        1 for usuario in usuarios
        if usuario.criado_em and usuario.criado_em.month == mes_atual
        and usuario.criado_em.year == ano_atual
    )
    return {
        'admin_name': admin_usuario.nome_completo if admin_usuario else 'Administrador',
        'admin_email': admin_usuario.email if admin_usuario else '',
        'usuarios': usuarios,
        'cameras': cameras,
        'alertas': alertas,
        'tokens': tokens,
        'relatorios': relatorios,
        'total_usuarios': len(usuarios),
        'usuarios_ativos': len(tokens),
        'total_cameras': len(cameras),
        'perfis_completos': perfis_completos,
        'perfis_pendentes': max(len(usuarios) - perfis_completos, 0),
        'receita': receita,
        'total_alertas': len(alertas),
        'novos_usuarios': novos_usuarios,
        'assinaturas_ativas': len(tokens),
    }


def ler_json(request):
    try:
        return json.loads(request.body or '{}')
    except json.JSONDecodeError:
        raise ValueError('JSON inválido')


def autenticado(view):
    @wraps(view)
    def wrapper(request, *args, **kwargs):
        valor = request.headers.get('Authorization', '')
        chave = valor.removeprefix('Bearer ').strip()
        try:
            request.usuario = TokenAcesso.objects.get(chave=chave).usuario
        except (TokenAcesso.DoesNotExist, TokenAcesso.MultipleObjectsReturned):
            return JsonResponse({'erro': 'Token ausente ou inválido'}, status=401)
        return view(request, *args, **kwargs)
    return wrapper


def administrador(view):
    @wraps(view)
    def wrapper(request, *args, **kwargs):
        usuario_id = request.session.get('admin_usuario_id')
        if not usuario_id:
            return redirect(f'/login/?next={request.path}')
        try:
            usuario = Usuario.objects.get(id=usuario_id, tipo_usuario='Admin')
        except (Usuario.DoesNotExist, ValueError):
            request.session.pop('admin_usuario_id', None)
            return redirect(f'/login/?next={request.path}')
        request.admin_usuario = usuario
        return view(request, *args, **kwargs)
    return wrapper


def destino_admin(valor):
    if valor and valor.startswith('/') and not valor.startswith('//'):
        return valor
    return '/'


def login_admin(request):
    if request.session.get('admin_usuario_id'):
        return redirect(destino_admin(request.GET.get('next')))

    erro = None
    if request.method == 'POST':
        email = request.POST.get('email', '').strip().lower()
        senha = request.POST.get('senha', '')
        try:
            usuario = Usuario.objects.get(email=email, tipo_usuario='Admin')
        except (Usuario.DoesNotExist, ValueError):
            usuario = None

        if usuario is not None and usuario.verificar_senha(senha):
            request.session['admin_usuario_id'] = str(usuario.id)
            return redirect(destino_admin(request.POST.get('next')))
        erro = 'E-mail ou senha inválidos.'

    return render(request, 'login.html', {
        'erro': erro,
        'next': request.GET.get('next', ''),
    })


def logout_admin(request):
    request.session.pop('admin_usuario_id', None)
    return redirect('/login/')


def cadastro_admin(request):
    erro = None
    if request.method == 'POST':
        nome = request.POST.get('nome', '').strip()
        email = request.POST.get('email', '').strip().lower()
        codigo = request.POST.get('codigo_admin', '')
        senha = request.POST.get('senha', '')
        confirmar_senha = request.POST.get('confirmar_senha', '')

        if not settings.ADMIN_REGISTRATION_CODE:
            erro = 'O cadastro administrativo está temporariamente indisponível.'
        elif codigo != settings.ADMIN_REGISTRATION_CODE:
            erro = 'Código de administrador inválido.'
        elif not nome or not email or len(senha) < 8:
            erro = 'Preencha os campos obrigatórios e use uma senha com no mínimo 8 caracteres.'
        elif senha != confirmar_senha:
            erro = 'As senhas não conferem.'
        elif Usuario.objects(email=email).first() is not None:
            erro = 'Este e-mail já está cadastrado.'
        else:
            usuario = Usuario(
                nome_completo=nome,
                email=email,
                tipo_usuario='Admin',
            )
            usuario.senha = make_password(senha)
            usuario.save()
            return redirect('/login/')

    return render(request, 'cadastro_admin.html', {'erro': erro})


@csrf_exempt
def login_api(request):
    if request.method != 'POST':
        return JsonResponse({'erro': 'Método não permitido'}, status=405)
    try:
        dados = ler_json(request)
        usuario = Usuario.objects.get(email=dados.get('email', '').lower())
    except (ValueError, Usuario.DoesNotExist):
        return JsonResponse({'erro': 'E-mail ou senha inválidos'}, status=401)
    if not usuario.verificar_senha(dados.get('senha', '')):
        return JsonResponse({'erro': 'E-mail ou senha inválidos'}, status=401)
    token = TokenAcesso.objects(usuario=usuario).first()
    if token is None:
        token = TokenAcesso(usuario=usuario, chave=secrets.token_urlsafe(32))
    else:
        token.chave = secrets.token_urlsafe(32)
    token.save()
    return JsonResponse({
        'token': token.chave,
        'usuario': serializar_usuario(usuario),
        'rastreamento': serializar_rastreamento(usuario),
    })


@csrf_exempt
def cadastro_api(request):
    if request.method != 'POST':
        return JsonResponse({'erro': 'Método não permitido'}, status=405)
    try:
        dados = ler_json(request)
        email = dados['email'].lower().strip()
        if not dados.get('nome_completo') or len(dados.get('senha', '')) < 6:
            raise ValueError('nome_completo e senha com no mínimo 6 caracteres são obrigatórios')
        if Usuario.objects(email=email).first() is not None:
            return JsonResponse({'erro': 'E-mail já cadastrado'}, status=409)
        usuario = Usuario(nome_completo=dados['nome_completo'], email=email,
                          ip_camera=dados.get('ip_camera'),
                          perfil_crianca=dados.get('perfil_crianca'))
        usuario.senha = make_password(dados['senha'])
        usuario.save()
    except (KeyError, ValueError):
        return JsonResponse({'erro': 'Dados de cadastro inválidos'}, status=400)
    return JsonResponse(serializar_usuario(usuario), status=201)


@autenticado
def perfil_api(request):
    return JsonResponse(serializar_usuario(request.usuario))


@autenticado
def alertas_api(request):
    alertas = Alerta.objects(responsavel=request.usuario).order_by('-timestamp')[:100]
    return JsonResponse([{'id': str(a.id), 'tipo_crise': a.tipo_crise,
                          'timestamp': a.timestamp.isoformat(),
                          'falso_positivo': a.falso_positivo,
                          'detalhes': a.detalhes} for a in alertas], safe=False)


@autenticado
@csrf_exempt
def camera_api(request):
    if request.method == 'GET':
        camera = Camera.objects(responsavel=request.usuario).first()
        if camera is None:
            return JsonResponse({'camera': None})
        return JsonResponse({
            'camera': {
                'id': str(camera.id),
                'endereco_ip': camera.endereco_ip,
                'nome_camara': camera.nome_camara,
                'localizacao': camera.localizacao,
                'modelo': camera.modelo,
            },
            'websocket': f'/ws/camera/{camera.id}/',
        })

    if request.method != 'POST':
        return JsonResponse({'erro': 'Método não permitido'}, status=405)
    try:
        dados = ler_json(request)
        endereco_ip = dados['endereco_ip'].strip()
        nome_camara = dados.get('nome_camara', 'Câmera principal').strip()
        if not endereco_ip or not nome_camara:
            raise ValueError
    except (KeyError, AttributeError, ValueError):
        return JsonResponse({'erro': 'endereco_ip e nome_camara são obrigatórios'}, status=400)

    camera = Camera.objects(responsavel=request.usuario).first()
    if camera is None:
        camera = Camera(responsavel=request.usuario, endereco_ip=endereco_ip,
                        nome_camara=nome_camara)
    else:
        camera.endereco_ip = endereco_ip
        camera.nome_camara = nome_camara
    camera.localizacao = dados.get('localizacao', camera.localizacao)
    camera.modelo = dados.get('modelo', camera.modelo)
    camera.save()

    request.usuario.ip_camera = endereco_ip
    request.usuario.save()
    return JsonResponse({
        'camera': {
            'id': str(camera.id),
            'endereco_ip': camera.endereco_ip,
            'nome_camara': camera.nome_camara,
        },
        'websocket': f'/ws/camera/{camera.id}/',
    }, status=201)


@method_decorator(administrador, name='dispatch')
class DashboardAdminView(View):
    def get(self, request):
        return render(request, 'admin_dashboard.html', contexto_admin(request.admin_usuario))


@method_decorator(administrador, name='dispatch')
class AdminTemplateView(View):
    template_name = None

    def get(self, request):
        return render(request, self.template_name, contexto_admin(request.admin_usuario))


class UsuariosAdminView(AdminTemplateView):
    template_name = 'usuarios.html'

    def get(self, request):
        return self.render_page(request)

    def post(self, request):
        nome = request.POST.get('nome_completo', '').strip()
        email = request.POST.get('email', '').strip().lower()
        senha = request.POST.get('senha', '')
        confirmar_senha = request.POST.get('confirmar_senha', '')
        ip_camera = request.POST.get('ip_camera', '').strip()
        erro = None

        if not nome or not email or len(senha) < 8:
            erro = 'Nome, e-mail e senha com no mínimo 8 caracteres são obrigatórios.'
        elif senha != confirmar_senha:
            erro = 'As senhas não conferem.'
        elif Usuario.objects(email=email).first() is not None:
            erro = 'Este e-mail já está cadastrado.'
        else:
            usuario = Usuario(
                nome_completo=nome,
                email=email,
                tipo_usuario='Admin',
                ip_camera=ip_camera or None,
            )
            usuario.senha = make_password(senha)
            try:
                usuario.save()
            except NotUniqueError:
                erro = 'Este e-mail já está cadastrado.'
            else:
                return redirect('usuarios')

        return self.render_page(request, {
            'erro_cadastro': erro,
            'form_cadastro': {
                'nome_completo': nome,
                'email': email,
                'ip_camera': ip_camera,
            },
        })

    def render_page(self, request, extra_context=None):
        contexto = contexto_admin(request.admin_usuario)
        busca = request.GET.get('q', '').strip()
        if busca:
            termo = busca.casefold()
            contexto['usuarios'] = [
                usuario for usuario in contexto['usuarios']
                if termo in (usuario.nome_completo or '').casefold()
                or termo in (usuario.email or '').casefold()
                or termo in (usuario.ip_camera or '').casefold()
            ]
        contexto['busca'] = busca
        if extra_context:
            contexto.update(extra_context)
        return render(request, self.template_name, contexto)


@method_decorator(administrador, name='dispatch')
class DetalharUsuarioAdminView(View):
    def get(self, request, usuario_id):
        try:
            usuario = Usuario.objects.get(id=usuario_id)
        except (Usuario.DoesNotExist, ValidationError, ValueError):
            return JsonResponse({'erro': 'Usuário não encontrado.'}, status=404)

        camera = Camera.objects(responsavel=usuario).first()
        return JsonResponse({
            'id': str(usuario.id),
            'nome_completo': usuario.nome_completo,
            'email': usuario.email,
            'tipo_usuario': usuario.tipo_usuario,
            'ip_camera': usuario.ip_camera or 'Não informado',
            'pagamento_conta': usuario.pagamento_conta or 'Não informado',
            'valor': usuario.valor,
            'criado_em': usuario.criado_em.strftime('%d/%m/%Y %H:%M') if usuario.criado_em else 'Não informado',
            'assinatura_ativa': TokenAcesso.objects(usuario=usuario).first() is not None,
            'camera': camera.nome_camara if camera else 'Não cadastrada',
            'perfil_crianca': serializar_perfil_crianca(usuario.perfil_crianca),
        })

    def post(self, request, usuario_id):
        try:
            usuario = Usuario.objects.get(id=usuario_id)
        except (Usuario.DoesNotExist, ValidationError, ValueError):
            return JsonResponse({'erro': 'Usuário não encontrado.'}, status=404)

        acao = request.POST.get('acao')
        if acao == 'excluir':
            if str(usuario.id) == str(request.admin_usuario.id):
                return JsonResponse({'erro': 'Você não pode excluir o administrador da sessão atual.'}, status=400)
            usuario.delete()
            return JsonResponse({'sucesso': True})

        if acao != 'editar':
            return JsonResponse({'erro': 'Ação inválida.'}, status=400)

        nome = request.POST.get('nome_completo', '').strip()
        email = request.POST.get('email', '').strip().lower()
        ip_camera = request.POST.get('ip_camera', '').strip()
        if not nome or not email:
            return JsonResponse({'erro': 'Nome e e-mail são obrigatórios.'}, status=400)
        if Usuario.objects(email=email, id__ne=usuario.id).first() is not None:
            return JsonResponse({'erro': 'Este e-mail já está cadastrado.'}, status=409)

        usuario.nome_completo = nome
        usuario.email = email
        usuario.ip_camera = ip_camera or None
        if any(chave.startswith('crianca_') for chave in request.POST):
            try:
                usuario.perfil_crianca = perfil_crianca_do_formulario(
                    request.POST, request.FILES, usuario.perfil_crianca,
                )
            except ValueError as erro_perfil:
                return JsonResponse({'erro': str(erro_perfil)}, status=400)
        try:
            usuario.save()
        except NotUniqueError:
            return JsonResponse({'erro': 'Este e-mail já está cadastrado.'}, status=409)
        return JsonResponse({'sucesso': True})


@method_decorator(administrador, name='dispatch')
class ExportarUsuariosAdminView(View):
    def get(self, request):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Usuários'
        worksheet.append(['ID', 'Nome completo', 'Tipo', 'E-mail', 'IP da câmera', 'Data de cadastro'])

        for celula in worksheet[1]:
            celula.font = Font(bold=True, color='FFFFFF')
            celula.fill = PatternFill('solid', fgColor='167B8A')

        for usuario in Usuario.objects.order_by('-criado_em'):
            worksheet.append([
                str(usuario.id),
                usuario.nome_completo or '',
                usuario.tipo_usuario or '',
                usuario.email or '',
                usuario.ip_camera or '',
                usuario.criado_em.strftime('%d/%m/%Y %H:%M') if usuario.criado_em else '',
            ])

        for coluna, largura in {'A': 26, 'B': 30, 'C': 18, 'D': 34, 'E': 20, 'F': 22}.items():
            worksheet.column_dimensions[coluna].width = largura
        worksheet.freeze_panes = 'A2'
        arquivo = BytesIO()
        workbook.save(arquivo)
        resposta = HttpResponse(
            arquivo.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resposta['Content-Disposition'] = 'attachment; filename="usuarios_neurai.xlsx"'
        return resposta


class AssinantesAdminView(AdminTemplateView):
    template_name = 'assinantes.html'

    def get(self, request):
        return self.render_page(request)

    def post(self, request):
        nome = request.POST.get('nome_completo', '').strip()
        email = request.POST.get('email', '').strip().lower()
        senha = request.POST.get('senha', '')
        confirmar_senha = request.POST.get('confirmar_senha', '')
        pagamento_conta = request.POST.get('pagamento_conta', '').strip()
        valor = request.POST.get('valor', '').strip()
        crianca_nome = request.POST.get('crianca_nome_completo', '').strip()
        erro = None

        if not nome or not email or len(senha) < 8:
            erro = 'Nome, e-mail e senha com no mínimo 8 caracteres são obrigatórios.'
        elif senha != confirmar_senha:
            erro = 'As senhas não conferem.'
        elif Usuario.objects(email=email).first() is not None:
            erro = 'Este e-mail já está cadastrado.'
        else:
            try:
                valor_numerico = float(valor.replace(',', '.')) if valor else None
            except ValueError:
                valor_numerico = None
                erro = 'O valor da assinatura é inválido.'

            if erro is None:
                try:
                    perfil = perfil_crianca_do_formulario(request.POST, request.FILES)
                    usuario = Usuario(
                        nome_completo=nome,
                        email=email,
                        tipo_usuario='Responsavel',
                        pagamento_conta=pagamento_conta or None,
                        valor=valor_numerico,
                        perfil_crianca=perfil,
                    )
                    usuario.senha = make_password(senha)
                    usuario.save()
                    TokenAcesso(usuario=usuario, chave=secrets.token_urlsafe(32)).save()
                except ValueError as erro_perfil:
                    erro = str(erro_perfil)
                except NotUniqueError:
                    erro = 'Este e-mail já está cadastrado.'
                else:
                    return redirect('assinantes')

        return self.render_page(request, {
            'erro_assinatura': erro,
            'form_assinatura': {
                'nome_completo': nome,
                'email': email,
                'pagamento_conta': pagamento_conta,
                'valor': valor,
                'crianca_nome_completo': crianca_nome,
            },
        })

    def render_page(self, request, extra_context=None):
        contexto = contexto_admin(request.admin_usuario)
        tokens = list(TokenAcesso.objects.order_by('-criado_em'))
        busca = request.GET.get('q', '').strip()
        termo = busca.casefold()
        assinantes = []
        for token in tokens:
            usuario = token.usuario
            if termo and not any(
                termo in (valor or '').casefold()
                for valor in (usuario.nome_completo, usuario.email, usuario.ip_camera)
            ):
                continue
            assinantes.append({
                'token': token,
                'usuario': usuario,
                'camera': Camera.objects(responsavel=usuario).first(),
            })

        mes_atual = datetime.utcnow().month
        ano_atual = datetime.utcnow().year
        novas_assinaturas = sum(
            1 for token in tokens
            if token.criado_em and token.criado_em.month == mes_atual
            and token.criado_em.year == ano_atual
        )
        contexto.update({
            'assinantes': assinantes,
            'busca': busca,
            'assinaturas_ativas': len(tokens),
            'novas_assinaturas': novas_assinaturas,
            'sem_assinatura': max(len(contexto['usuarios']) - len(tokens), 0),
        })
        if extra_context:
            contexto.update(extra_context)
        return render(request, self.template_name, contexto)


@method_decorator(administrador, name='dispatch')
class ExportarAssinantesAdminView(View):
    def get(self, request):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Assinaturas'
        worksheet.append(['ID do usuário', 'Assinante', 'E-mail', 'Câmera', 'Data da assinatura', 'Status'])

        for celula in worksheet[1]:
            celula.font = Font(bold=True, color='FFFFFF')
            celula.fill = PatternFill('solid', fgColor='167B8A')

        for token in TokenAcesso.objects.order_by('-criado_em'):
            usuario = token.usuario
            camera = Camera.objects(responsavel=usuario).first()
            worksheet.append([
                str(usuario.id),
                usuario.nome_completo or '',
                usuario.email or '',
                camera.nome_camara if camera else 'Não cadastrada',
                token.criado_em.strftime('%d/%m/%Y %H:%M') if token.criado_em else '',
                'Ativa',
            ])

        for coluna, largura in {'A': 26, 'B': 30, 'C': 34, 'D': 24, 'E': 24, 'F': 14}.items():
            worksheet.column_dimensions[coluna].width = largura
        worksheet.freeze_panes = 'A2'
        arquivo = BytesIO()
        workbook.save(arquivo)
        resposta = HttpResponse(
            arquivo.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resposta['Content-Disposition'] = 'attachment; filename="assinaturas_neurai.xlsx"'
        return resposta


class RelatoriosAdminView(AdminTemplateView):
    template_name = 'relatorios.html'


class FinanceiroAdminView(AdminTemplateView):
    template_name = 'financeiro.html'


class CadastrosAdminView(AdminTemplateView):
    template_name = 'cadastros.html'